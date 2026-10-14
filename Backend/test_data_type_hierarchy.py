from app.ontology.namespace import FDO
from app.ontology.owl_builder import build_fdo_ontology
from app.services.comparison_service import evaluate_data_match
from app.services.data_type_hierarchy import data_type_relation
from app.services.data_type_hierarchy import (
    canonical_data_type_code,
    is_known_data_type,
)
from app.services.fdo_decision_service import (
    FDODecisionInput,
    evaluate_fdo_decision,
)
from app.services.ancillary_baseline import is_ancillary_declaration
from app.ontology.rdf_instance_builder import build_rdf_graph
from app.ontology.classes import ExpectedData, DeclaredData
from app.ontology.namespace import FDO
from rdflib import OWL, RDF, RDFS
import sqlite3
from pathlib import Path

from openpyxl import load_workbook


def test_equal_data_types_are_exact():
    assert data_type_relation(
        "DATA.LOCATION.PRECISE",
        "DATA.LOCATION.PRECISE",
    ) == "EXACT"
    assert evaluate_data_match(
        "DATA.LOCATION.PRECISE",
        "DATA.LOCATION.PRECISE",
    ) == "EXACT_MATCH"


def test_broader_evidence_matches_expected_child():
    assert data_type_relation(
        "DATA.LOCATION.PRECISE",
        "DATA.LOCATION",
    ) == "BROADER_EVIDENCE"
    assert evaluate_data_match(
        "DATA.LOCATION.PRECISE",
        "DATA.LOCATION",
    ) == "HIERARCHICAL_MATCH"


def test_narrower_evidence_matches_expected_parent():
    assert data_type_relation(
        "DATA.LOCATION",
        "DATA.LOCATION.PRECISE",
    ) == "NARROWER_EVIDENCE"


def test_sibling_subtypes_do_not_match():
    assert data_type_relation(
        "DATA.LOCATION.PRECISE",
        "DATA.LOCATION.APPROXIMATE",
    ) is None


def test_legacy_comm_namespace_resolves_to_canonical_comm_namespace():
    assert data_type_relation(
        "DATA.COMM.MESSAGE_CONTENT",
        "DATA.COM.MESSAGE_CONTENT",
    ) == "EXACT"


def test_coarse_personal_info_does_not_match_an_unrelated_account_subtree():
    assert data_type_relation(
        "DATA.ACCOUNT.GOOGLE",
        "DATA.ACCOUNT.PERSONAL_INFO",
    ) is None


def test_legacy_location_label_uses_the_canonical_hierarchy():
    assert data_type_relation(
        "DATA.LOCATION.PRECISE",
        "precise_location",
    ) == "EXACT"


def test_data_comm_namespace_is_aliased_to_data_comm():
    assert canonical_data_type_code(
        "DATA.COM.MESSAGE_CONTENT"
    ) == "DATA.COMM.MESSAGE_CONTENT"
    assert data_type_relation(
        "DATA.COMM.MESSAGE_CONTENT",
        "DATA.COM.MESSAGE_CONTENT",
    ) == "EXACT"


def test_added_type_codes_have_sensible_parents():
    expected_parents = {
        "DATA.DEVICE.SENSOR": "DATA.DEVICE",
        "DATA.HEALTH.PERSONAL_HEALTH": "DATA.HEALTH",
        "DATA.HEALTH.VITALS": "DATA.HEALTH",
        "DATA.MEDIA.AUDIO_VIDEO": "DATA.MEDIA",
        "DATA.USAGE.ACTIVITY_LOGS": "DATA.USAGE",
        "DATA.USAGE.TRIP_DETAILS": "DATA.USAGE",
    }
    from app.services.data_type_hierarchy import DATA_TYPE_PARENTS

    for data_type, parent in expected_parents.items():
        assert DATA_TYPE_PARENTS[data_type] == parent


def test_stage3_and_database_data_types_resolve_to_hierarchy_or_alias():
    repository_root = Path(__file__).resolve().parents[1]
    workbook = load_workbook(
        repository_root
        / "DATA"
        / "FDO-Privacy_Final_Stage3_Dataset_CLEANED.xlsx",
        read_only=True,
        data_only=True,
    )
    worksheet = workbook["Expected Data"]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    type_index = headers.index("Data Type")
    data_types = {
        str(row[type_index]).strip()
        for row in rows
        if row[type_index] is not None
    }
    workbook.close()

    connection = sqlite3.connect(
        f"file:{repository_root / 'Backend' / 'fdo_privacy.db'}?mode=ro",
        uri=True,
    )
    try:
        for table_name in ("expected_mappings", "policy_evidence"):
            data_types.update(
                row[0]
                for row in connection.execute(
                    f"SELECT DISTINCT data_type FROM {table_name} "
                    "WHERE data_type IS NOT NULL"
                )
            )
    finally:
        connection.close()

    assert data_types
    unresolved = sorted(
        data_type
        for data_type in data_types
        if not is_known_data_type(data_type)
    )
    assert unresolved == []


def test_related_location_types_do_not_fire_rule_5():
    cases = [
        ("DATA.LOCATION.PRECISE", "DATA.LOCATION"),
        ("DATA.LOCATION", "DATA.LOCATION.PRECISE"),
    ]

    for expected_type, evidence_type in cases:
        assert evaluate_data_match(expected_type, evidence_type) == (
            "HIERARCHICAL_MATCH"
        )
        decision = evaluate_fdo_decision(
            FDODecisionInput(
                expected_id="expected-location",
                application_id="app-location",
                expected_data_type=expected_type,
                expected_necessity="REQUIRED",
                expected_purpose="navigation",
                declared_evidence_id="evidence-location",
                declared_data_type=evidence_type,
                declared_purpose="navigation",
                data_match="HIERARCHICAL_MATCH",
                purpose_compatibility="COMPATIBLE",
                functional_relationship="RELATED",
                evidence_status="EXTRACTED",
            )
        )
        assert decision.fdo_status == "NOT_FDO"
        assert "differs from" not in decision.decision_reason


def test_rule_3_broader_evidence_has_lower_confidence_and_reason():
    decision = evaluate_fdo_decision(
        FDODecisionInput(
            expected_id="expected-precise",
            application_id="app-location",
            expected_data_type="DATA.LOCATION.PRECISE",
            expected_necessity="REQUIRED",
            expected_purpose="navigation",
            declared_evidence_id="evidence-location",
            declared_data_type="DATA.LOCATION",
            declared_purpose="navigation",
            data_match="HIERARCHICAL_MATCH",
            purpose_compatibility="COMPATIBLE",
            functional_relationship="RELATED",
            evidence_status="EXTRACTED",
            data_match_relation="BROADER_EVIDENCE",
        )
    )

    assert decision.fdo_status == "NOT_FDO"
    assert decision.decision_confidence == 0.75
    assert "matched via broader type" in decision.decision_reason
    assert decision.data_match_relation == "BROADER_EVIDENCE"


def test_data_type_hierarchy_is_declared_in_ontology():
    graph = build_fdo_ontology()

    assert (
        FDO.DataType_DATA_LOCATION_PRECISE,
        RDF.type,
        OWL.Class,
    ) in graph
    assert (
        FDO.DataType_DATA_LOCATION_PRECISE,
        RDFS.subClassOf,
        FDO.DataType_DATA_LOCATION,
    ) in graph
    assert (
        FDO.DataType_DATA_LOCATION,
        RDFS.subClassOf,
        FDO.DataType_DATA,
    ) in graph
    assert (
        FDO.DataType_DATA,
        RDFS.subClassOf,
        FDO.DataType,
    ) in graph


def test_expected_and_declared_nodes_link_literal_and_canonical_type_class():
    graph = build_rdf_graph(
        expected_data=[
            ExpectedData(
                id="EXP-COM",
                name="message",
                data_type="DATA.COM.MESSAGE_CONTENT",
                purpose="messaging",
            )
        ],
        declared_data=[
            DeclaredData(
                id="EV-COM",
                name="message",
                data_type="DATA.COM.MESSAGE_CONTENT",
                purpose="messaging",
                extraction_status="EXTRACTED",
            )
        ],
    )
    expected_uri = FDO["resource/expected-data/EXP-COM"]
    declared_uri = FDO["resource/declared-data/EV-COM"]
    expected_class = FDO.DataType_DATA_COMM_MESSAGE_CONTENT
    declared_class = expected_class

    assert (expected_uri, FDO.dataType, None) in graph
    assert (expected_uri, FDO.dataTypeClass, expected_class) in graph
    assert (declared_uri, FDO.dataType, None) in graph
    assert (declared_uri, FDO.dataTypeClass, declared_class) in graph


def test_ancillary_baseline_requires_both_type_and_purpose_category():
    assert is_ancillary_declaration(
        "DATA.DEVICE.ID",
        "security verification",
    )
    assert is_ancillary_declaration(
        "DATA.ACCOUNT.USER_ID",
        "service operation",
    )
    assert is_ancillary_declaration(
        "DATA.AUTH.SECRET_KEY",
        "account security",
    )
    assert not is_ancillary_declaration(
        "DATA.HEALTH.RECORD",
        "security",
    )
    assert not is_ancillary_declaration(
        "DATA.DEVICE.ID",
        "marketing",
    )
