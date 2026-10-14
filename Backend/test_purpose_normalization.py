import sqlite3
from pathlib import Path

from openpyxl import load_workbook

from app.services.fdo_decision_service import (
    FDODecisionInput,
    evaluate_fdo_decision,
)
from app.services.comparison_service import (
    normalize_purpose,
    classify_purpose,
    classify_purposes,
    evaluate_purpose_compatibility,
)


def test_normalize_basic_text():
    assert normalize_purpose("Navigation") == "navigation"


def test_normalize_whitespace():
    assert normalize_purpose("  navigation   ") == "navigation"


def test_normalize_hyphen():
    assert normalize_purpose("route-planning") == "route planning"


def test_normalize_underscore():
    assert normalize_purpose("route_planning") == "route planning"


def test_normalize_slash():
    assert normalize_purpose("navigation/route planning") == (
        "navigation route planning"
    )


def test_normalize_repeated_spaces():
    assert normalize_purpose("route     planning") == "route planning"


def test_empty_purpose():
    assert normalize_purpose("") == ""


def test_none_purpose():
    assert normalize_purpose(None) == ""


def test_exact_match_after_normalization():
    assert evaluate_purpose_compatibility(
        "Navigation",
        " navigation ",
    ) == "COMPATIBLE"


def test_hyphen_normalization_match():
    assert evaluate_purpose_compatibility(
        "route planning",
        "route-planning",
    ) == "COMPATIBLE"


def test_different_purposes_remain_incompatible():
    assert evaluate_purpose_compatibility(
        "navigation",
        "advertising",
    ) == "INCOMPATIBLE"


def test_missing_declared_purpose():
    assert evaluate_purpose_compatibility(
        "navigation",
        None,
    ) == "NO_EVIDENCE"

def test_classify_navigation():
    assert classify_purpose("navigation") == "NAVIGATION"


def test_classify_directions():
    assert classify_purpose("directions") == "NAVIGATION"


def test_classify_route_planning():
    assert classify_purpose("route planning") == "NAVIGATION"


def test_classify_nearby_search():
    assert classify_purpose("nearby search") == "NAVIGATION"


def test_classify_combined_navigation_purpose():
    assert classify_purpose(
        "directions and nearby search"
    ) == "NAVIGATION"


def test_classify_advertising():
    assert classify_purpose(
        "targeted advertising"
    ) == "ADVERTISING"


def test_classify_payment():
    assert classify_purpose(
        "processing transactions"
    ) == "PAYMENT_PROCESSING"


def test_classify_account_management():
    assert classify_purpose(
        "account management"
    ) == "ACCOUNT_MANAGEMENT"


def test_classify_unknown_purpose():
    assert classify_purpose(
        "something completely unrelated"
    ) is None


def test_classify_empty_purpose():
    assert classify_purpose("") is None


def test_classify_none_purpose():
    assert classify_purpose(None) is None

def test_semantic_navigation_directions():
    assert evaluate_purpose_compatibility(
        "navigation",
        "directions",
    ) == "COMPATIBLE"


def test_semantic_navigation_route_planning():
    assert evaluate_purpose_compatibility(
        "navigation",
        "route planning",
    ) == "COMPATIBLE"


def test_semantic_navigation_nearby_search():
    assert evaluate_purpose_compatibility(
        "navigation",
        "nearby search",
    ) == "COMPATIBLE"


def test_semantic_navigation_combined_purpose():
    assert evaluate_purpose_compatibility(
        "navigation",
        "directions and nearby search",
    ) == "COMPATIBLE"


def test_navigation_advertising_incompatible():
    assert evaluate_purpose_compatibility(
        "navigation",
        "targeted advertising",
    ) == "INCOMPATIBLE"


def test_payment_transaction_compatibility():
    assert evaluate_purpose_compatibility(
        "payment processing",
        "processing transactions",
    ) == "COMPATIBLE"


def test_account_authentication_compatibility():
    assert evaluate_purpose_compatibility(
        "account management",
        "authentication",
    ) == "COMPATIBLE"


def test_unknown_purpose_is_conservative():
    assert evaluate_purpose_compatibility(
        "navigation",
        "completely unknown purpose",
    ) == "UNKNOWN"


def test_known_unknown_purpose_is_unknown():
    assert evaluate_purpose_compatibility(
        "navigation",
        "an unclassified purpose",
    ) == "UNKNOWN"


def test_unknown_unknown_purpose_is_unknown():
    assert evaluate_purpose_compatibility(
        "an unclassified purpose",
        "another unclassified purpose",
    ) == "UNKNOWN"


def test_known_incompatible_purposes_remain_incompatible():
    assert evaluate_purpose_compatibility(
        "navigation",
        "targeted advertising",
    ) == "INCOMPATIBLE"


def test_unknown_purpose_routes_decision_to_ambiguous():
    decision = evaluate_fdo_decision(
        FDODecisionInput(
            expected_id="expected-1",
            application_id="app-1",
            expected_data_type="DATA.LOCATION",
            expected_necessity="REQUIRED",
            expected_purpose="navigation",
            declared_evidence_id="evidence-1",
            declared_data_type="DATA.LOCATION",
            declared_purpose="unclassified",
            data_match="EXACT_MATCH",
            purpose_compatibility="UNKNOWN",
            functional_relationship="RELATED",
            evidence_status="EXTRACTED",
        )
    )

    assert decision.fdo_status == "AMBIGUOUS"


def test_every_expected_purpose_is_classified():
    repository_root = Path(__file__).resolve().parents[1]
    workbook = load_workbook(
        repository_root
        / "DATA"
        / "FDO-Privacy_Final_Stage3_Dataset_CLEANED.xlsx",
        read_only=True,
        data_only=True,
    )
    worksheet = workbook["Expected Data"]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    purpose_index = headers.index("Purpose")
    purposes = {
        str(row[purpose_index]).strip()
        for row in rows
        if row[purpose_index] is not None
    }

    connection = sqlite3.connect(
        repository_root / "Backend" / "fdo_privacy.db"
    )
    try:
        purposes.update(
            row[0]
            for row in connection.execute(
                "SELECT DISTINCT expected_purpose "
                "FROM expected_mappings "
                "WHERE expected_purpose IS NOT NULL"
            )
        )
    finally:
        connection.close()

    assert purposes
    assert all(classify_purpose(purpose) for purpose in purposes)


def test_expected_purpose_categories_include_domain_specific_phrases():
    expected_categories = {
        "File storage/sync": "FILE_STORAGE_SYNC",
        "Device control": "DEVICE_CONTROL",
        "Activity tracking": "ACTIVITY_TRACKING",
        "Fitness tracking": "FITNESS_TRACKING",
        "Search/localization": "SEARCH_LOCATION",
        "Security/verification": "SECURITY",
    }
    assert {
        purpose: classify_purpose(purpose)
        for purpose in expected_categories
    } == expected_categories

def test_classify_purposes_single_category():
    result = classify_purposes("directions")

    assert result == ["NAVIGATION"]


def test_classify_purposes_multiple_categories():
    result = classify_purposes(
        "directions and targeted advertising"
    )

    assert result == ["NAVIGATION", "ADVERTISING"]


def test_classify_purposes_multiple_phrases_same_category():
    result = classify_purposes(
        "directions and nearby search"
    )

    assert result == ["NAVIGATION"]


def test_classify_purposes_navigation_and_analytics():
    result = classify_purposes(
        "directions and usage analysis"
    )

    assert result == ["NAVIGATION", "ANALYTICS"]


def test_classify_purposes_unknown():
    result = classify_purposes(
        "quantum optimization"
    )

    assert result == []


def test_classify_purposes_empty():
    result = classify_purposes("")

    assert result == []


def test_classify_purposes_none():
    result = classify_purposes(None)

    assert result == []

def test_multi_purpose_compatible():
    result = evaluate_purpose_compatibility(
        "navigation",
        "directions and targeted advertising",
    )

    assert result == "COMPATIBLE"


def test_multi_purpose_incompatible():
    result = evaluate_purpose_compatibility(
        "navigation",
        "targeted advertising and marketing",
    )

    assert result == "INCOMPATIBLE"


def test_multi_purpose_navigation_and_analytics():
    result = evaluate_purpose_compatibility(
        "navigation",
        "directions and usage analysis",
    )

    assert result == "COMPATIBLE"


def test_multi_purpose_payment_and_marketing():
    result = evaluate_purpose_compatibility(
        "payments",
        "targeted advertising and marketing",
    )

    assert result == "INCOMPATIBLE"